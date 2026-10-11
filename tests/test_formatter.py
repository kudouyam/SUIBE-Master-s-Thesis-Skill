# Run: pytest -q tests/test_formatter.py
from pathlib import Path
import json
import subprocess
import sys
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
TOOL=ROOT/'scripts'/'format_thesis.py'


def format_tmp(doc,*,args=()):
    with tempfile.TemporaryDirectory() as tmp:
        inp=Path(tmp)/'in.docx';out=Path(tmp)/'out.docx'
        doc.save(inp)
        run=subprocess.run([sys.executable,str(TOOL),str(inp),'--out',str(out),*args],
                           text=True,capture_output=True)
        assert run.returncode==0,(run.stdout,run.stderr)
        data=json.loads((Path(tmp)/'in_格式修正报告.json').read_text(encoding='utf-8'))
        yield Document(out),data,inp,out


def test_mixed_levels_toc_and_body():
    d=Document();d.add_paragraph('论文题目：财务治理研究')
    d.add_paragraph('摘要');d.add_paragraph('中文摘要。'*100)
    d.add_paragraph('Abstract');d.add_paragraph('An abstract in English.')
    d.add_paragraph('目录')
    d.add_paragraph('1 第一章………………1')
    d.add_paragraph('1 绪论')
    d.add_paragraph('这是一段包含中文与 English 的普通正文。')
    d.add_paragraph('1.1 背景');d.add_paragraph('1.1.1 贡献')
    d.add_paragraph('1.1.1.1 样本');d.add_paragraph('（1）数据处理')
    for output,data,_,_ in format_tmp(d):
        ps=output.paragraphs
        assert ps[6].style.name!='Heading 1', 'TOC must not turn into Heading 1'
        assert ps[7].style.name=='Heading 1'
        assert ps[8].runs[0].font.size.pt==12
        assert [x['level'] for x in data['headings']]==[1,2,3,4,5]
        assert ps[7].text=='1\u3000绪论'
        assert ps[7].paragraph_format.page_break_before is True
        ind=ps[8]._p.xpath('./w:pPr/w:ind')[0]
        assert ind.get(qn('w:firstLineChars'))=='200'
        assert ps[6].runs[-1].font.size.pt==12, 'TOC page numbers always 12pt'
        assert not any(x['severity']=='severe' for x in data['issues'])


def test_captions_three_line_white_and_integrity():
    d=Document();d.add_paragraph('1 第一章');d.add_paragraph('正文。')
    t=d.add_table(rows=3,cols=2);t.style='Light Shading Accent 1'
    t.cell(0,0).text='year';t.cell(0,1).text='revenue'
    t.cell(1,0).text='2024';t.cell(1,1).text='2'
    t.cell(2,0).text='2025';t.cell(2,1).text='3'
    d.add_paragraph('表 1.1 收入增长统计')
    for out,data,_,fname in format_tmp(d):
        # Table caption must be before table, with no content removed.
        b=out.element.body
        body=[x for x in b.iterchildren() if x.tag in (qn('w:p'),qn('w:tbl'))]
        idx=next(i for i,x in enumerate(body) if x.tag==qn('w:tbl'))
        assert '表 1.1' in ''.join(x.text or '' for x in body[idx-1].xpath('.//w:t'))
        tbl=out.tables[0]
        borders=tbl._tbl.tblPr.find(qn('w:tblBorders'))
        assert borders.find(qn('w:insideV')).get(qn('w:val'))=='nil'
        assert tbl.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
        for row in tbl.rows:
            for c in row.cells:
                assert c._tc.tcPr.find(qn('w:shd')).get(qn('w:fill'))=='FFFFFF'
        assert not any(x['severity']=='severe' for x in data['issues'])
        assert 'word/document.xml' in ZipFile(fname).namelist()


def test_complex_merged_table_is_preserved():
    d=Document();d.add_paragraph('1 第一章')
    t=d.add_table(rows=2,cols=2);t.cell(0,0).merge(t.cell(0,1))
    for _,data,_,_ in format_tmp(d):
        assert any(x['category']=='complex_table' for x in data['issues'])
        assert data['summary']['changes_by_type'].get('复杂表格保留')==1


def test_abstract_body_and_cover_preserved():
    d=Document();d.add_paragraph('作者姓名：某某');d.add_paragraph('论文题目：质量控制研究')
    d.add_paragraph('摘要');p=d.add_paragraph('论文研究了合理的内部控制。');p.runs[0].font.size=Pt(21)
    d.add_paragraph('Abstract');q=d.add_paragraph('Internal control research.');q.runs[0].font.size=Pt(18)
    d.add_paragraph('目录');d.add_paragraph('1 一章……1');d.add_paragraph('1 绪论')
    for out, data, _, _ in format_tmp(d):
        assert out.paragraphs[0].text=='作者姓名：某某'
        assert out.paragraphs[3].runs[0].font.size.pt==12
        assert out.paragraphs[5].runs[0].font.size.pt==12
        assert len(out.sections)==2
        assert '质量控制研究' in out.sections[1].header.paragraphs[0].text
        assert not out.sections[0].header.paragraphs[0].text.strip()
        assert any(x['category']=='odd_page' for x in data['issues'])


def test_chinese_heading_hierarchy():
    d=Document();d.add_paragraph('一、背景')
    d.add_paragraph('（一）理论基础');d.add_paragraph('1. 财务研究')
    d.add_paragraph('（1）研究意义');d.add_paragraph('① 第一要点')
    d.add_paragraph('二、结论')
    for out,data,_,_ in format_tmp(d):
        assert data['summary']['number_mode']=='chinese'
        assert [x['level'] for x in data['headings'][:5]]==[1,2,3,4,5]
        assert any(x['category']=='numbering' for x in data['issues'])


def test_semantic_body_shading_preserved_even_in_unmerged_table():
    d=Document();d.add_paragraph('1 财务研究')
    t=d.add_table(rows=3,cols=2);t.style='Light Shading Accent 1'
    t.cell(0,0).text='指标';t.cell(0,1).text='数值'
    t.cell(1,0).text='增长';t.cell(1,1).text='高风险'
    t.cell(2,0).text='利润';t.cell(2,1).text='100'
    p=t.cell(1,1)._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'FF6666');p.append(shd)
    for out,data,_,_ in format_tmp(d):
        colors=out.tables[0].cell(1,1)._tc.xpath('./w:tcPr/w:shd')
        assert colors and colors[-1].get(qn('w:fill'))=='FF6666'
        assert any(i['category']=='semantic_table_color' for i in data['issues'])
        assert data['summary']['diagnostics']['table_semantic_colors_preserved']==1


def test_list_style_is_not_misidentified_as_chapter():
    d=Document();d.add_paragraph('1 第一章')
    d.add_paragraph('这是主要的正文内容。')
    d.add_paragraph('1. 准备材料',style='List Number')
    d.add_paragraph('2. 开展调查',style='List Number')
    for out,data,_,_ in format_tmp(d):
        assert [x['level'] for x in data['headings']]==[1]
        assert out.paragraphs[2].style.name=='List Number'
        assert out.paragraphs[3].style.name=='List Number'


def test_inline_figure_caption_is_safely_moved_below_picture():
    from PIL import Image
    from docx.shared import Inches
    with tempfile.TemporaryDirectory() as temp:
        image=Path(temp)/'image.png'
        Image.new('RGB',(22,22),'white').save(image)
        d=Document();d.add_paragraph('1 绪论')
        d.add_paragraph('图 1-1 模拟测试图')
        d.add_paragraph().add_run().add_picture(str(image),width=Inches(.5))
        d.add_paragraph('资料来源：自制测试样例')
        for out,data,_,_ in format_tmp(d):
            blocks=[e for e in out.element.body.iterchildren() if e.tag in (qn('w:p'),qn('w:tbl'))]
            caption_idx=next(i for i,b in enumerate(blocks) if '图 1-1' in ''.join(b.xpath('.//w:t/text()')))
            assert bool(blocks[caption_idx-1].xpath('.//w:drawing'))
            assert data['summary']['diagnostics']['figure_caption_auto_relocated']==1
            assert not any(x['severity']=='severe' for x in data['issues'])


def test_figures_numbering_duplicates_and_chapter_conflicts_reported():
    d=Document();d.add_paragraph('1 理论基础')
    d.add_paragraph('图 1-1 理论框架')
    d.add_paragraph('图 1-1 另一框架')
    d.add_paragraph('图 2-3 非本章图')
    for out,data,_,_ in format_tmp(d):
        categories={x['category'] for x in data['issues']}
        assert 'caption_duplicate' in categories
        assert 'caption_chapter_mismatch' in categories
        assert 'caption_sequence' in categories


def test_fourth_heading_does_not_inherit_word_default_italics():
    d=Document()
    d.add_paragraph('1 第一章');d.add_paragraph('1.1 第一节')
    d.add_paragraph('1.1.1 第一小节');d.add_paragraph('1.1.1.1 四级标题')
    for out, data,_,_ in format_tmp(d):
        p=out.paragraphs[3]
        assert p.style.name=='Heading 4'
        assert p.style.font.italic is False
        assert all(r.italic is False for r in p.runs if r.text)
