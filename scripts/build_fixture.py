"""Create a deliberately messy test thesis to smoke test SUIBE formatter."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path(__file__).resolve().parent.parent/'examples'/'格式混乱的测试论文.docx'
d=Document()
d.add_paragraph('上海对外经贸大学')
d.add_paragraph('硕士学位论文')
d.add_paragraph('论文题目：公司财务风险与治理研究')
d.add_paragraph('作者姓名：测试用户')
d.add_paragraph('学位论文原创性声明')
d.add_paragraph('本人所呈交的学位论文是我在导师指导下完成的。')
d.add_paragraph('摘要')
p=d.add_paragraph('本文研究公司财务风险问题。'*40)
p.runs[0].font.size=Pt(17)
d.add_paragraph('【关键词】企业；风险；治理')
d.add_paragraph('Abstract')
d.add_paragraph('This paper describes corporate financial risk management.'*4)
d.add_paragraph('【Keywords】Risk; Corporate Governance')
d.add_paragraph('目录')
d.add_paragraph('1 绪论…………………………………………1')
d.add_paragraph('1.1 研究背景…………………………………2')
d.add_paragraph('1.1.1 研究方法………………………………3')
d.add_paragraph('1 绪论')
d.add_paragraph('近年来，全球经济发展变化很大，企业需要更高质量的财务分析方法。')
d.add_paragraph('1.1 研究背景')
d.add_paragraph('本文采用案例分析法，讨论了企业内部治理与现金流之间的联系。')
d.add_paragraph('1.1.1 研究方法')
d.add_paragraph('公司现金流是一个重要的观察指标。公司现金流管理需要结合不同维度的数据。')
d.add_paragraph('1.1.1.1 数据来源')
d.add_paragraph('本项目数据来自公开披露的财务信息。')
d.add_paragraph('（1）进一步讨论')
d.add_paragraph('该细节需要结合实际状况进行进一步研究。')
d.add_paragraph('2 实证分析')
d.add_paragraph('以下从不同角度观察数据的分布与差异。')
t=d.add_table(rows=3, cols=3)
t.style='Light Shading Accent 1'
t.cell(0,0).text='指标';t.cell(0,1).text='2020年';t.cell(0,2).text='2021年'
t.cell(1,0).text='收入';t.cell(1,1).text='12';t.cell(1,2).text='13'
t.cell(2,0).text='利润';t.cell(2,1).text='5';t.cell(2,2).text='6'
d.add_paragraph('表 2.1 样本公司经营表现')  # deliberately incorrect position and separator
p=d.add_paragraph()
p.add_run('图形占位文字（真实图形使用 Word Drawing 对象）')
d.add_paragraph('图 2-1 样本结构图')
d.add_paragraph('参考文献')
d.add_paragraph('[1]张三.经济学研究[J].某刊,2020,1(1):1-10.')
d.add_paragraph('[2]John Smith. Data Analysis[M]. New York: Press, 2021.')
d.add_paragraph('致谢')
d.add_paragraph('衷心感谢研究期间给予帮助的导师与同学。')
d.save(out)
print(out)
