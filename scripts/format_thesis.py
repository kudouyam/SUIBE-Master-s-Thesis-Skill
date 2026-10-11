#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SUIBE 2020 thesis formatter, conservative OOXML preserving edition.

This tool is a deterministic first-pass formatter, not a substitute for Word
visual inspection or for the university-supplied cover/declaration templates.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import sys
import traceback
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.run import Run
from copy import deepcopy
from lxml import etree

VERSION = '1.1.0'
NAME = '上海对外经贸大学硕士学位论文内容和格式要求（2020年修订）'
SOURCE = '规范第3-7页、同套格式范例第3-8页'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
SIZES = {1: 15, 2: 14, 3: 12, 4: 12, 5: 12}
HAN_NUM = '一二三四五六七八九十百千〇零两'

RE_TABLE = re.compile(r'^\s*表\s*[\d一二三四五六七八九十]+\s*[-—－.．]\s*\d+\s*\S+', re.I)
RE_FIGURE = re.compile(r'^\s*图\s*[\d一二三四五六七八九十]+\s*[-—－.．]\s*\d+\s*\S+', re.I)
RE_REF = re.compile(r'^\s*(\[\d{1,4}\]|［\d+］|\d{1,4}\.)\s*\S.{12,}', re.I)
RE_CHAPTER = re.compile(r'^第\s*[\d一二三四五六七八九十百]+\s*章\s*\S+')
RE_SECTION = re.compile(r'^第\s*[\d一二三四五六七八九十百]+\s*节\s*\S+')
RE_DECIMAL = re.compile(r'^(\d{1,3}(?:\.\d{1,3}){1,3})\s*[.．、]?\s*(\S.*)$')
RE_ARABIC_MAIN = re.compile(r'^(\d{1,3})\s*(?:[.．、]|\s)\s*(\S.*)$')
RE_CN_MAIN = re.compile(r'^([一二三四五六七八九十]+)、\s*(\S.*)$')
RE_CN_SECOND = re.compile(r'^[（(]([一二三四五六七八九十]+)[）)]\s*(\S.*)$')
RE_CN_FIG = re.compile(r'^[（(](\d{1,3})[）)]\s*(\S.*)$')
RE_CIRCLE = re.compile(r'^[①②③④⑤⑥⑦⑧⑨⑩]\s*(\S.*)$')
TITLE_WORDS = {'摘要', '中文摘要', 'Abstract', 'ABSTRACT', '目录', '目 录',
               '参考文献', '附录', '致谢', '后记', '结论', '绪论', '引言',
               '学位论文原创性声明', '学位论文使用授权声明'}
SPECIALS = {'摘要':'abstract_cn', '中文摘要':'abstract_cn', 'Abstract':'abstract_en',
            'ABSTRACT':'abstract_en', '目录':'toc_title', '目 录':'toc_title',
            '参考文献':'references_title','附录':'appendix_title',
            '致谢':'ack_title', '后记':'ack_title',
            '学位论文原创性声明':'declaration','学位论文使用授权声明':'declaration'}


def el_first(parent, tag):
    ch = parent.find(qn('w:' + tag))
    if ch is None:
        ch = OxmlElement('w:' + tag)
        parent.append(ch)
    return ch


def set_font(run_or_style, *, size:float, east_asia='宋体', latin='Times New Roman', bold=None, black=False):
    """Set only font face/size without flattening existing runs or losing emphasis."""
    if hasattr(run_or_style, '_r'):
        r = run_or_style._r
        if any(r.xpath('.//w:drawing|.//w:pict|.//w:object|.//w:fldChar|.//w:instrText|.//w:sym|.//m:oMath')):
            return False
        rpr = r.get_or_add_rPr()
    else:
        rpr = run_or_style._element.get_or_add_rPr()
    fonts = el_first(rpr, 'rFonts')
    for a, val in [('ascii', latin), ('hAnsi', latin), ('eastAsia', east_asia), ('cs', latin)]:
        fonts.set(qn('w:' + a), val)
    for a in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + a), None)
    for nm in ('sz', 'szCs'):
        t = el_first(rpr, nm)
        t.set(qn('w:val'), str(int(round(size * 2))))
    # Explicitly return text character tracking to the default.
    for child in list(rpr):
        if child.tag in (qn('w:spacing'), qn('w:w')):
            rpr.remove(child)
    if black:
        color=el_first(rpr,'color')
        color.set(qn('w:val'),'000000')
        color.attrib.pop(qn('w:themeColor'),None)
    if bold is not None:
        b = el_first(rpr, 'b')
        b.set(qn('w:val'), '1' if bold else '0')
    return True


def set_p_format(p, *, first_line_chars=None, keep_next=False, page_break=False,
                 align=None, line=20, reset_left=True):
    pf = p.paragraph_format
    pf.line_spacing = Pt(line)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True
    pf.keep_with_next = bool(keep_next)
    if page_break:
        pf.page_break_before = True
    if align is not None:
        p.alignment = align
    if first_line_chars is not None:
        ppr = p._p.get_or_add_pPr()
        ind = el_first(ppr, 'ind')
        if reset_left:
            for k in ['left', 'leftChars', 'start', 'startChars', 'hanging', 'hangingChars', 'firstLine']:
                ind.attrib.pop(qn('w:' + k), None)
        if first_line_chars == 0:
            ind.attrib.pop(qn('w:firstLineChars'), None)
        else:
            ind.set(qn('w:firstLineChars'), str(first_line_chars * 100))


def has_complex_xml(element):
    return bool(element.xpath('.//w:ins|.//w:del|.//w:sdt|.//w:object|.//w:pict|.//w:txbxContent|.//w:altChunk'))


def get_outline(p):
    lev = p._p.xpath('./w:pPr/w:outlineLvl')
    if lev:
        try: return int(lev[0].get(qn('w:val'))) + 1
        except Exception: pass
    return None


def get_list_level(p):
    levels = p._p.xpath('./w:pPr/w:numPr/w:ilvl')
    if not levels: return None
    try: return int(levels[0].get(qn('w:val'))) + 1
    except Exception: return None


def bold_dominant(p):
    runs = [r for r in p.runs if r.text.strip()]
    if not runs: return False
    b = [r for r in runs if r.bold is True or (r.bold is None and p.style.font.bold is True)]
    return len(b) / len(runs) >= 0.7


def has_toc_pattern(s, p):
    style = p.style.name.lower()
    if re.search(r'\btoc\s*[1-9]\b|目录\s*[1-9一二三]', style): return True
    if p._p.xpath('.//w:instrText'):
        if any('TOC ' in ''.join(x.itertext()).upper() for x in p._p.xpath('.//w:instrText')):
            return True
    return bool(re.search(r'(?:\t|…{3,}|\.{4,}|·{4,})\s*\d+\s*$', s))


def determine_mode(paras):
    cn = sum(bool(RE_CN_MAIN.match(p.text.strip())) for p in paras)
    decimal = sum(bool(RE_DECIMAL.match(p.text.strip())) for p in paras)
    ch = sum(bool(RE_CHAPTER.match(p.text.strip())) for p in paras)
    return 'chinese' if cn >= 2 and cn > decimal and cn >= ch else 'decimal'


def simple_structure(doc):
    return {
       'paragraphs':len(doc.paragraphs), 'tables':len(doc.tables),
       'paragraph_texts':Counter(p.text for p in doc.paragraphs),
       'table_texts':Counter(c.text for t in doc.tables for row in t.rows for c in row.cells),
    }


def media_fingerprint(path):
    with ZipFile(path) as z:
        return {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}


def package_xml_stats(path):
    with ZipFile(path) as z:
        names = set(z.namelist())
        xml = z.read('word/document.xml')
        root = etree.fromstring(xml)
        out = {'pictures':len(root.xpath('.//w:drawing|.//w:pict',namespaces=NS)),
               'math_objects':len(root.xpath('.//m:oMath',namespaces=NS)),
               'complex_control_count':len(root.xpath('.//w:sdt|.//w:txbxContent|.//w:object',namespaces=NS)),
               'tracked_change_nodes':len(root.xpath('.//w:ins|.//w:del',namespaces=NS)),
               'has_footnotes':'word/footnotes.xml' in names,
               'has_comments':'word/comments.xml' in names,
               'has_toc_field':bool(re.search(rb'\bTOC\b',xml)),
               'has_section_breaks':len(root.xpath('.//w:sectPr', namespaces=NS))}
        return out


class Work:
    def __init__(self, doc, infile, *, title=None, table_mode='simple', headers='safe'):
        self.doc = doc
        self.infile = infile
        self.title = title.strip() if title else None
        self.table_mode = table_mode
        self.headers = headers
        self.mode = determine_mode(doc.paragraphs)
        self.issues = []
        self.changes = []
        self.headings = []
        self.captions = []
        self.sections = []
        self.counts = Counter()
        self.diagnostics = {'table_semantic_colors_preserved': 0, 'figure_caption_auto_relocated': 0}
        self.section_context = {}
        self.toc_indexes = set()
        self.frontmatter_indexes = set()
        self.recognized = {}
        self._collect_context()

    def issue(self, level, kind, description, location='', action='人工核查'):
        self.issues.append({'severity':level,'category':kind,'location':location,
                            'detail':description,'recommended_action':action})

    def change(self, target, location, before, after, origin):
        self.changes.append({'target':target,'location':location,'before':before,
                             'after':after,'evidence':origin})
        self.counts[target]+=1

    def _collect_context(self):
        markers = {}
        for i,p in enumerate(self.doc.paragraphs):
            s=p.text.strip()
            kind=SPECIALS.get(s)
            if kind: markers.setdefault(kind,[]).append(i)
            if '论文题目' in s and not self.title:
                m=re.search(r'论文题目\s*[:：]\s*(.+)',s)
                if m and not re.fullmatch(r'(?:X+|x+|论文题目|题目|待填|\s+)',m.group(1).strip()):
                    self.title=m.group(1).strip()
        if not self.title:
            metadata = self.doc.core_properties.title.strip() if self.doc.core_properties.title else ''
            if metadata and metadata not in {'文档','Document','论文','XXX'}:
                self.title=metadata
        # Mark definite front matter by position up to first actual chapter and deduced TOC area.
        front_keys = ('abstract_cn','abstract_en','toc_title','declaration')
        pos=[i for k in front_keys for i in markers.get(k,[])]
        toc_idx=markers.get('toc_title',[None])[0]
        toc_started = toc_idx is not None
        if pos:
            first_body=None
            for j,p in enumerate(self.doc.paragraphs):
                if j <= min(pos): continue
                s=p.text.strip()
                if (RE_CHAPTER.match(s) or RE_CN_MAIN.match(s) or RE_DECIMAL.match(s)
                      or RE_ARABIC_MAIN.match(s)) and not has_toc_pattern(s,p):
                    # TOC-looking entries are usually short with page numbers; if TOC region,
                    # require supporting long body immediately afterward to escape it.
                    near=' '.join(x.text.strip() for x in self.doc.paragraphs[j+1:j+4])
                    next_text=self.doc.paragraphs[j+1].text.strip() if j+1<len(self.doc.paragraphs) else ''
                    next_is_prose=bool(len(next_text)>16 and re.search(r'[。；，,.]',next_text)
                                       and not has_toc_pattern(next_text,self.doc.paragraphs[j+1])) if j+1<len(self.doc.paragraphs) else False
                    strong=bool(next_is_prose or (re.search(r'[。；，,.]',near) and len(near)>55)
                                or p.paragraph_format.page_break_before)
                    if toc_idx is None or j<toc_idx or strong:
                        first_body=j;break
            if first_body is not None:
                self.frontmatter_indexes.update(range(first_body))
            elif toc_idx is not None:
                self.issue('review','frontmatter','检测到目录但无法可靠定位正文起点；只处理能明确识别的区域。','目录','核对正文首页及目录条目')
        if toc_started:
            for j in range(toc_idx+1,len(self.doc.paragraphs)):
                if j not in self.frontmatter_indexes and j>toc_idx+2: break
                s=self.doc.paragraphs[j].text.strip()
                if has_toc_pattern(s,self.doc.paragraphs[j]) or j in self.frontmatter_indexes:
                    self.toc_indexes.add(j)
        self.markers=markers
        if not self.title:
            self.issue('review','header','无法从文档中可靠获取真实论文题目；不会把“XXX”等占位文字当成页眉。', '页眉','用 --title "真实论文标题" 重新运行')

    def classify(self,p,i):
        s=p.text.strip()
        if not s: return ('blank',0,'none')
        kind=SPECIALS.get(s)
        if kind:return (kind,0,'certain')
        if i in self.toc_indexes or has_toc_pattern(s,p):return ('toc_entry',0,'certain')
        if RE_TABLE.match(s):return ('table_caption',0,'certain')
        if RE_FIGURE.match(s):return ('figure_caption',0,'certain')
        if re.fullmatch(r'【?关键词】?.*|【?Keywords】?.*',s,re.I):return ('keywords',0,'certain')
        if i in self.frontmatter_indexes:
            if any(x<=i for x in self.markers.get('abstract_en',[])) and 'abstract' in self.section_context.get('last',''):
                return ('body_en',0,'normal')
            return ('frontmatter',0,'normal')
        if s in {'绪论','引言','结论'}:return ('heading',1,'medium')
        if RE_REF.match(s) and self.section_context.get('references',False):return ('reference_item',0,'certain')
        sty=p.style.name
        # Numbered-list body paragraphs are not headings merely because they
        # visually start with '1.': protect ordinary List/List Paragraph styles.
        if re.search(r'(?:^|\b)(?:List(?: Paragraph)?|列表|项目符号|编号)(?:\s|$|[0-9])',sty,re.I):
            if not re.search(r'(?:Heading|标题)\s*[1-5]\b',sty,re.I):
                return ('list_body',0,'style/list')
        m=re.search(r'(?:Heading|标题|标题\s*级)\s*([1-5])\b', sty, re.I)
        if m:
            return ('heading',int(m.group(1)),'high/style')
        ol=get_outline(p)
        if ol and 1<=ol<=5:
            return ('heading',ol,'high/outline')
        # Longer prose with punctuation should not be identified as a heading.
        if len(s)>100 or (len(s)>38 and s.endswith(('。','；',';','，',',','.'))):
            return ('body',0,'normal')
        if RE_CHAPTER.match(s): return ('heading',1,'high/chapter')
        if RE_SECTION.match(s): return ('heading',2,'high/section')
        dec=RE_DECIMAL.match(s)
        if dec:
            return ('heading',min(dec.group(1).count('.')+1,5),'high/number')
        if self.mode=='chinese':
            if RE_CN_MAIN.match(s):return ('heading',1,'high/chinese')
            if RE_CN_SECOND.match(s):return ('heading',2,'high/chinese')
            if RE_ARABIC_MAIN.match(s):return ('heading',3,'medium/chinese')
            if RE_CN_FIG.match(s):return ('heading',4,'medium/chinese')
            if RE_CIRCLE.match(s):return ('heading',5,'medium/chinese')
        else:
            if RE_CN_MAIN.match(s):return ('heading',1,'medium/chinese')
            if RE_CN_SECOND.match(s):return ('heading',2,'medium/chinese')
            if RE_ARABIC_MAIN.match(s):return ('heading',1,'high/number')
            if RE_CN_FIG.match(s) and bold_dominant(p):return ('heading',5,'medium/number')
        if i>0 and get_list_level(p) is not None and len(s)<45 and bold_dominant(p):
            lev=get_list_level(p)
            self.issue('review','ambiguous_heading',f'自动编号加粗段落可能是第{lev}级标题，也可能是列表；没有强制归级：{s[:75]}',f'正文第{i+1}段','核对其多级列表及编号')
            return ('preserve',0,'uncertain/list')
        if len(s)<35 and bold_dominant(p) and s not in TITLE_WORDS:
            self.issue('review','ambiguous_heading',f'无明确编号但像标题，保持现有格式：{s[:60]}',f'正文第{i+1}段','人工确认标题级别')
            return ('preserve',0,'uncertain/bold')
        if RE_REF.match(s) and len(s)>30:return ('reference_item',0,'normal')
        return ('body',0,'normal')

    def ensure_body_section(self):
        """Safely separate the main-body story from preliminary pages if possible."""
        if len(self.doc.sections)!=1:return
        structured_front=bool(self.markers.get('toc_title') or self.markers.get('abstract_cn')
                              or self.markers.get('declaration'))
        if not structured_front:return
        one=[h for h in self.headings if h['level']==1]
        if not one:return
        first=one[0]['paragraph']-1
        if first<=0 or (self.frontmatter_indexes and first-1 not in self.frontmatter_indexes):
            self.issue('review','section_boundary','无法可靠确认正文与前置部分的分界；未创建正文分节。',
                       '正文首页','Word 手工设置下一奇数页分节')
            return
        # Existing complex revision stories should not be mechanically re-sectioned.
        if self.doc.element.body.xpath('.//w:ins|.//w:del|.//w:sdt'):
            self.issue('review','section_boundary','文档含追踪修订或内容控件，不自动移动分节语义。',
                       '正文首页','在 Word 校对修订后设置奇数页分节')
            return
        old=self.doc.element.body.sectPr
        if old is None:
            self.issue('review','section_boundary','缺少主文档节属性，未强制分节。','正文首页')
            return
        prev=self.doc.paragraphs[first-1]
        ppr=prev._p.get_or_add_pPr()
        if ppr.find(qn('w:sectPr')) is not None:
            return
        front_sect=deepcopy(old)
        oldtype=front_sect.find(qn('w:type'))
        if oldtype is not None:front_sect.remove(oldtype)
        ppr.append(front_sect)
        bodytype=el_first(old,'type')
        bodytype.set(qn('w:val'),'oddPage')
        pn=el_first(old,'pgNumType')
        pn.set(qn('w:start'),'1')
        self.change('正文与前置部分分节',f'正文第{first+1}段之前',
                    '一个 Word 分节（前置与正文混合）',
                    '在正文前插入分节边界；正文从下一右页/奇数页开始，页码从1起算',
                    '规范第3页、范例第6页')
        self.issue('review','odd_page','已设置正文从下一奇数页起排且页码从1开始，但 Word 具体分页、空白页数量与分节继承关系仍需渲染核对。',
                   '正文首章','检查前置页面、正文首页及分节后的页眉页码')

    def set_section_layout(self):
        for i,sec in enumerate(self.doc.sections):
            before=f'{sec.top_margin.mm:.1f}/{sec.bottom_margin.mm:.1f}/{sec.left_margin.mm:.1f}/{sec.right_margin.mm:.1f}mm'
            sec.page_width=Mm(210);sec.page_height=Mm(297)
            sec.left_margin=Mm(35); sec.right_margin=Mm(25)
            sec.top_margin=Mm(37); sec.bottom_margin=Mm(30)
            self.change('A4页面及边距',f'第{i+1}节',before,'A4, 上37/下30/左35/右25mm','规范第7页')
        self.issue('info','source_conflict','规范第7页打印区“155×230nm”与明确边距不吻合，采用明确边距（内容宽150mm、高230mm）。','全稿','无需按155mm改动')

    def set_styles(self):
        st=self.doc.styles
        normal=st['Normal']
        set_font(normal,size=12)
        normal.paragraph_format.line_spacing=Pt(20)
        normal.paragraph_format.space_before=Pt(0)
        normal.paragraph_format.space_after=Pt(0)
        for lvl in range(1,6):
            key=f'Heading {lvl}'
            s=st[key]
            set_font(s,size=SIZES[lvl],bold=True,black=True)
            s.font.italic=False  # Word built-in Heading 4 often inherits italics
            s.font.underline=False
            s.paragraph_format.line_spacing=Pt(20)
            s.paragraph_format.space_before=Pt(0)
            s.paragraph_format.space_after=Pt(0)
            s.paragraph_format.keep_with_next=True
            s.paragraph_format.widow_control=True
            s.paragraph_format.left_indent=Pt(0)
            s.paragraph_format.first_line_indent=Pt(0)
        for level in (4,5):
            self.issue('info','unspecified_rule',f'第{level}级标题使用宋体小四12pt加粗作为保守建议值，校规未明确字号。',f'标题{level}级','如学院另有细则，以学院规定为准')

    def set_toc_style(self,p,level):
        set_p_format(p,first_line_chars=0,align=WD_ALIGN_PARAGRAPH.LEFT)
        for r in p.runs:
            set_font(r,size=SIZES[level],bold=True,black=True)
        # TOC page number must be small-four 12pt in every hierarchy.
        runs=[r for r in p.runs if r.text]
        if not runs:return
        r=runs[-1]
        match=re.search(r'(\d{1,5})\s*$',r.text)
        if not match:return
        if r._r.xpath('.//w:fldChar|.//w:instrText|.//w:drawing|.//w:pict|.//w:object'):
            self.issue('review','toc_page_font','目录页码由域/复杂run生成，不能安全按片段调字号。',p.text[:50],
                       'Word 更新目录后检查页码为宋体小四')
            return
        left=r.text[:match.start(1)]; tail=r.text[match.start(1):]
        if not left:
            set_font(r,size=12,east_asia='宋体',latin='Times New Roman',bold=False,black=True)
        elif len(r._r.xpath('./w:t'))==1 and all(ch.tag in (qn('w:rPr'),qn('w:t')) for ch in r._r):
            new_node=deepcopy(r._r)
            r._r.addnext(new_node)
            r.text=left
            added=Run(new_node,p)
            added.text=tail
            set_font(added,size=12,east_asia='宋体',latin='Times New Roman',bold=False,black=True)
        else:
            self.issue('review','toc_page_font','目录条目页码与标题共用复杂run，未强行拆分以防域损坏。',p.text[:55],
                       '检查右侧页码的宋体小四12pt')

    def normalize_heading_gap(self,p,loc):
        """Correct only plain, single-run numbered-heading gaps."""
        nonempty=[r for r in p.runs if r.text.strip()]
        if len(nonempty)!=1:return
        run=nonempty[0]
        if len(run._r.xpath('./w:t'))!=1 or run._r.xpath('./w:drawing|./w:fldChar|./w:instrText|./w:sym'):
            return
        old=run.text
        pattern=r'^(\s*)(\d{1,3}(?:\.\d{1,3}){0,3}[.．、]?|第[一二三四五六七八九十百\d]+[章节]|[一二三四五六七八九十]+、|[（(][一二三四五六七八九十\d]+[）)])([ \t\u3000]*)(\S.*)$'
        m=re.match(pattern,old)
        if not m:return
        if m.group(3)=='　':return
        run.text=m.group(1)+m.group(2)+'　'+m.group(4)
        self.change('标题编号间距',loc,repr(m.group(3)),'编号与标题之间1个全角汉字空白（U+3000）','规范第4页')

    def process_paragraphs(self):
        cn_ab=None
        latest_heading_level=None
        for i,p in enumerate(self.doc.paragraphs):
            s=p.text.strip()
            if not s: continue
            kind,level,confidence=self.classify(p,i)
            # Only a short '（1）...' following a recognized L4 heading may be L5
            # in the school's decimal scheme. Never blanket-convert numeric lists.
            if (kind=='body' and self.mode=='decimal' and latest_heading_level==4
                    and RE_CN_FIG.match(s) and len(s)<=48):
                kind,level,confidence='heading',5,'medium/context-L4'
            loc=f'正文第{i+1}段（{s[:32]}）'
            self.recognized[i]={'kind':kind,'level':level,'confidence':confidence,'text':s[:125]}
            if kind=='abstract_cn':cn_ab='cn'
            if kind=='abstract_en':cn_ab='en'
            if kind=='toc_title':cn_ab=None
            if kind in ['references_title','ack_title','appendix_title']:
                self.section_context['references']=kind=='references_title'
                cn_ab=None
            if kind=='heading':
                if level>5:
                    self.issue('review','ambiguous_heading',f'识别到超过五级标题：{s[:75]}',loc)
                    continue
                if has_complex_xml(p._p):
                    self.issue('review','complex_content','复杂标题含有修订/文本框等结构，跳过直接修改。',loc)
                    continue
                before=p.style.name
                p.style=self.doc.styles[f'Heading {level}']
                set_p_format(p,first_line_chars=0,keep_next=True,
                             page_break=(level==1),align=WD_ALIGN_PARAGRAPH.LEFT)
                for r in p.runs:
                    if set_font(r,size=SIZES[level],bold=True,black=True):
                        r.italic=False
                        r.underline=False
                self.normalize_heading_gap(p,loc)
                self.headings.append({'paragraph':i+1,'level':level,'text':p.text.strip(),'confidence':confidence})
                latest_heading_level=level
                self.change(f'{level}级标题',loc,before,f'宋体{SIZES[level]}pt加粗；左对齐；顶格；固定20pt',
                            '范例第6页；规范第4页')
                if level==1:self.section_context['references']=False
                if (RE_CN_MAIN.match(s) or RE_CN_SECOND.match(s) or
                    self.mode=='chinese' and (RE_ARABIC_MAIN.match(s) or RE_CN_FIG.match(s))):
                    self.issue('review','numbering',f'原始标题编号采用中文层级或混排，已归级排版，尚未重写编号或全文交叉引用：{s[:70]}',loc,'按规范1/1.1/1.1.1重编并核对所有引用')
            elif kind in ('table_caption','figure_caption'):
                set_p_format(p,first_line_chars=0,keep_next=(kind=='table_caption'),
                             align=WD_ALIGN_PARAGRAPH.CENTER)
                for r in p.runs:set_font(r,size=9,bold=True,black=True)
                self.captions.append({'paragraph':i+1,'kind':kind,'text':s})
                self.change('图表题',loc,'原始图表题格式','宋体小五9pt加粗；已保留原编号与文字',
                            '规范第4-5页')
                if re.search(r'\d+[.．]\d+',s):
                    self.issue('review','caption_number',f'范例曾用小数点编号，但正式规范示例为“章-序号”；未自动改号以免破坏交叉引用：{s[:70]}',loc,'如需统一，连同正文引用一起核对')
            elif kind in ('toc_title','references_title','ack_title','appendix_title'):
                set_p_format(p,first_line_chars=0,keep_next=True,page_break=kind in ('toc_title','references_title','ack_title','appendix_title'),
                             align=WD_ALIGN_PARAGRAPH.CENTER)
                for r in p.runs:set_font(r,size=15,bold=True,black=True)
                self.change('前后置标题',loc,'原格式','宋体小三15pt居中（目录、参考文献、附录、致谢）',
                            '范例第5、7、8页；部分为建议')
            elif kind in ('abstract_cn','abstract_en'):
                set_p_format(p,first_line_chars=0,keep_next=True,
                             page_break=(kind in ('abstract_cn','abstract_en')),align=WD_ALIGN_PARAGRAPH.CENTER)
                for r in p.runs:set_font(r,size=16,bold=None,black=True)
                self.change('摘要标题',loc,'原格式','宋体三号16pt；英文标题字体按范例保守处理','规范第3页')
            elif kind=='toc_entry':
                m=RE_DECIMAL.match(s)
                level=m.group(1).count('.')+1 if m else 1
                if level<=3:self.set_toc_style(p,level)
                else:self.issue('review','toc','目录项疑似包含四级以上标题，正式规范要求目录至三级。',loc)
                self.change('目录条目',loc,'原格式','按目录层级1~3设置宋体加粗固定20pt；页码需在Word更新','规范第3页')
            elif kind=='frontmatter' and cn_ab in ('cn','en'):
                if has_complex_xml(p._p):
                    self.issue('review','abstract_object','摘要含复杂嵌入对象，未直接修改。',loc)
                    continue
                set_p_format(p,first_line_chars=2,align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                for r in p.runs:set_font(r,size=12,east_asia='宋体',latin='Times New Roman',black=True)
                self.change('中外文摘要正文',loc,'原有格式','小四12pt、中文宋体/英文Times New Roman，固定20pt','规范第3页')
            elif kind=='list_body':
                # Keep Word list numbering/hanging-indent semantics intact while
                # standardizing only text face/size/line spacing.
                set_p_format(p,first_line_chars=None,align=None)
                for r in p.runs:
                    set_font(r,size=12,east_asia='宋体',latin='Times New Roman')
                self.change('编号列表文字',loc,'保留列表结构与编号',
                            '只统一宋体/Times New Roman小四12pt及固定20pt行距',
                            '规范第3页正文样式；保护Word自动编号结构')
            elif kind in ('frontmatter','declaration','preserve','blank'):
                # Official cover/declarations should never be reconstructed by this script.
                self.counts['保留前置或不确定段落']+=1
            elif kind=='keywords':
                set_p_format(p,first_line_chars=0)
                for r in p.runs:
                    set_font(r,size=12,east_asia='宋体',latin='Times New Roman')
                self.change('关键词',loc,'原字号','宋体/Times New Roman小四（建议值）','范例第3-4页')
            else:
                if has_complex_xml(p._p):
                    self.issue('review','complex_content','正文段落含文本框/修订/对象等，保留复杂段落结构与直接格式。',loc)
                    continue
                if kind == 'body_en' or cn_ab=='en':
                    font_ea='宋体'; font_l='Times New Roman'
                elif kind=='reference_item':
                    font_ea='宋体'; font_l='Times New Roman'
                elif kind=='body' and self.section_context.get('references',False):
                    font_ea='宋体';font_l='Times New Roman'
                elif kind=='body' and s in ['论文作者签名','导师签名']:
                    continue
                else:
                    font_ea='宋体';font_l='Times New Roman'
                if not (kind=='reference_item' or self.section_context.get('references',False)):
                    set_p_format(p,first_line_chars=2,align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                else:
                    set_p_format(p,first_line_chars=0)
                if kind=='body' and self._in_ack(i): font_ea='楷体'
                for r in p.runs: set_font(r,size=12,east_asia=font_ea,latin=font_l)
                self.change('正文/文献/致谢文字',loc,'各原始格式','中文'+font_ea+'、英文Times New Roman小四12pt；固定20pt',
                            '规范第3-4、7页；范例第7-8页')
        if self.mode=='chinese':
            self.issue('review','numbering','检测到原稿以中文序号为主，已调整格式层级；未擅自重编号或篡改“见第X节”等交叉引用。',
                       '全文标题','最终提交前按学校阿拉伯编号体系统一并更新引用')

    def _in_ack(self,i):
        ack=self.markers.get('ack_title',[])
        if not ack:return False
        return i>ack[-1]

    def process_tables(self):
        for ix,t in enumerate(self.doc.tables,1):
            loc=f'第{ix}个表格'
            rows=len(t.rows);cols=len(t.columns)
            complex_table=any(cell._tc.xpath('.//w:vMerge|.//w:gridSpan|.//w:sdt|.//w:drawing|.//w:pict|.//m:oMath')
                              for row in t.rows for cell in row.cells)
            nested=any(cell.tables for row in t.rows for cell in row.cells)
            if complex_table or nested or cols>12 or rows<2:
                self.issue('review','complex_table',f'{rows}行×{cols}列，含合并/复杂对象/嵌套表或列数过多；保留内部边框与颜色，避免数据含义被改变。',
                           loc,'手工核对三线表、重复表头、续表及来源')
                self.counts['复杂表格保留']+=1
                continue
            # Cells in the body (not the heading row) may encode meaning by color:
            # heat maps, negative/positive highlights, status or annotated significance.
            # These are not equivalent to a decorative single-colored header.
            body_colors = set()
            for row in t.rows[1:]:
                for cell in row.cells:
                    for shd in cell._tc.xpath('./w:tcPr/w:shd'):
                        fill=(shd.get(qn('w:fill')) or '').upper()
                        if fill and fill not in {'AUTO','FFFFFF','FFF','NONE'}:
                            body_colors.add(fill)
            if body_colors:
                self.issue('review','semantic_table_color',
                           f'检测到表格正文单元格有非白色填充 {sorted(body_colors)[:8]}；可能传递分组/显著性/热力信息，已保留整张表的原有颜色与边框。',
                           loc,'人工确认颜色是否有含义，再决定是否转换三线表')
                self.counts['有意义底色表格保留']+=1
                self.diagnostics['table_semantic_colors_preserved']+=1
                continue
            if self.table_mode == 'preserve':
                self.issue('info','table_preserve','选择表格保守模式：保留原始边框/底色；只执行整体文档格式化。',loc)
                continue
            try:
                t.alignment=WD_TABLE_ALIGNMENT.CENTER
                try:t.style=self.doc.styles['Normal Table']
                except KeyError:pass
                tblpr=t._tbl.tblPr
                borders=el_first(tblpr,'tblBorders')
                for name in ['top','bottom','left','right','insideH','insideV']:
                    tag=el_first(borders,name)
                    if name in ('top','bottom'):
                        tag.set(qn('w:val'),'single');tag.set(qn('w:sz'),'8')
                        tag.set(qn('w:color'),'000000')
                    else:tag.set(qn('w:val'),'nil')
                for row in t.rows:
                    for cell in row.cells:
                        tcpr=cell._tc.get_or_add_tcPr()
                        for old in list(tcpr):
                            if old.tag in (qn('w:tcBorders'),qn('w:shd')):
                                tcpr.remove(old)
                        shd=OxmlElement('w:shd');shd.set(qn('w:val'),'clear')
                        shd.set(qn('w:fill'),'FFFFFF')
                        tcpr.append(shd)
                # Draw rule after header by applying bottom border only to first-row cells.
                for cell in t.rows[0].cells:
                    tcpr=cell._tc.get_or_add_tcPr()
                    cbd=el_first(tcpr,'tcBorders')
                    bottom=el_first(cbd,'bottom')
                    bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'8')
                    bottom.set(qn('w:color'),'000000')
                head=t.rows[0]._tr.get_or_add_trPr()
                el_first(head,'tblHeader').set(qn('w:val'),'true')
                self.change('简单三线表',loc,f'{rows}行×{cols}列，原始样式不一',
                            '黑色上/表头下/下横线；无竖线、白底；重复表头',
                            '规范第4页建议＋范例第6页白底')
            except Exception as e:
                self.issue('review','table_error',f'该表修改失败：{type(e).__name__}: {e}',loc,'恢复原稿该表并单独修复')

    def check_caption_locations(self):
        # Blocks only within Word main document story; no fabricated visual positions.
        body=self.doc.element.body
        blocks=[e for e in body.iterchildren() if e.tag in (qn('w:p'),qn('w:tbl'))]
        for j,block in enumerate(blocks):
            if block.tag!=qn('w:p'):continue
            s=''.join(el.text or '' for el in block.xpath('.//w:t')).strip()
            if RE_TABLE.match(s):
                next_block=blocks[j+1] if j+1<len(blocks) else None
                prev_block=blocks[j-1] if j>0 else None
                if next_block is not None and next_block.tag==qn('w:tbl'):
                    pass
                elif prev_block is not None and prev_block.tag==qn('w:tbl'):
                    # Safely reposition whole caption paragraph before the table only
                    # if it is an unambiguous *immediately adjacent* caption.
                    prev_block.addprevious(block)
                    self.change('表题位置','表题：'+s[:30],'表格之后','移到紧邻表格正上方','规范第4页')
                else:
                    self.issue('review','caption_location',f'表题与数据表未紧邻，无法确定归属：{s[:50]}',
                               '表题段落','人工确认表题位于目标表格正上方')
            if RE_FIGURE.match(s):
                prev_block=blocks[j-1] if j>0 else None
                next_block=blocks[j+1] if j+1<len(blocks) else None
                preceding=prev_block is not None and bool(prev_block.xpath('.//w:drawing|.//w:pict'))
                following=next_block is not None and bool(next_block.xpath('.//w:drawing|.//w:pict'))
                if following and not preceding:
                    # One immediately following pure inline-image paragraph is safe to
                    # reorder as a complete OOXML paragraph; floating shapes must stay.
                    is_inline=bool(next_block.xpath('.//wp:inline'))
                    is_anchor=bool(next_block.xpath('.//wp:anchor'))
                    has_nonempty_text=bool(''.join(next_block.xpath('.//w:t/text()')).strip())
                    if is_inline and not is_anchor and not has_nonempty_text and not has_complex_xml(next_block):
                        next_block.addnext(block)
                        self.diagnostics['figure_caption_auto_relocated']+=1
                        self.change('图题位置','图题：'+s[:30],
                            '紧邻行内图片的上方','移至该行内图片下方（完整段落挪动，图像关系保留）','规范第5页')
                    else:
                        self.issue('review','figure_caption_location',f'图题疑似在图片上方，但图片为浮动/复合结构，不自动移动：{s[:60]}',
                                   '图题段落','人工检查图片锚点并将图题调整到图下方')
                if not following and not preceding:
                    self.issue('info','figure_caption_location',f'无法只用段落相邻关系确认图题是否在图片下方：{s[:60]}',
                               '图题段落','渲染后核对')
        if self.doc.tables and not self.captions:
            self.issue('review','missing_captions','存在表格但未识别到符合“表 1-1 XXX”等形式的表题。',
                       '所有表格','人工补齐表号、表题及资料来源（不得让程序凭空编造来源）')

    def footnotes(self):
        # Settings for Word-generated footnote numbering by page.
        settings=self.doc.settings._element
        fp=el_first(settings,'footnotePr')
        nr=el_first(fp,'numRestart')
        nr.set(qn('w:val'),'eachPage')
        if package_xml_stats(self.infile)['has_footnotes']:
            self.issue('info','footnotes','已有脚注 XML；保存后将设为小五9pt、单倍行距，并设 Word 脚注按页重编号。',
                       '所有脚注','重新在 Word 更新并验证分页编号')
        else:
            self.issue('info','footnotes','未在包内检测到现有脚注；设置了今后脚注每页重新编号的 Word 属性。','全文')

    def format_headers_footers(self):
        structured_front=bool(self.markers.get('toc_title') or self.markers.get('abstract_cn') or self.markers.get('declaration'))
        secnum=len(self.doc.sections)
        if secnum==1 and structured_front:
            self.issue('review','header_footer','文档只含一个 Word 分节并含封面/摘要/目录：若直接设置页眉页码可能污染封面，故不自动写入；需在正文起点插入合适分节并核查奇数页起始。',
                       '正文首章与前置内容分界','Word 分节后设置正文页眉标题与居中 PAGE 域')
            return
        if secnum>1:
            self.issue('review','header_footer','存在多个分节。自动处理正文候选分节，但每一节首页、连续编号/重新编号需人工核对。',
                       '节分隔符','逐节检查首页不同、链接到前一节等')
        for j,sec in enumerate(self.doc.sections):
            if j==0 and secnum>1:continue
            if self.headers=='preserve':continue
            if sec.header.is_linked_to_previous and j>0:
                sec.header.is_linked_to_previous=False
            if sec.footer.is_linked_to_previous and j>0:
                sec.footer.is_linked_to_previous=False
            h=sec.header
            hp=h.paragraphs[0]
            if self.title and not hp.text.strip() and not hp._p.xpath('.//w:fldChar'):
                hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=hp.add_run(self.title)
                set_font(r,size=9)
                self.change('页眉',f'第{j+1}节','空','论文真实题目居中','规范第3页；字号9pt为非校规默认')
            elif hp.text.strip() and self.title and hp.text.strip()!=self.title:
                self.issue('review','header','现有页眉不同于论文标题，保留原值并请求核对。',f'第{j+1}节页眉','校正为论文标题')
            fp=sec.footer.paragraphs[0]
            fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if not fp.text.strip() and not fp._p.xpath('.//w:fldChar'):
                r=fp.add_run()
                begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
                ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' PAGE '
                separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate')
                cached=OxmlElement('w:t');cached.text='1'
                end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
                for item in [begin,ins,separate,cached,end]: r._r.append(item)
                self.change('页脚页码',f'第{j+1}节','空','居中的 Word PAGE 字段','规范第3页')
            else:
                self.issue('info','footer_preserve','存在页脚文字/页码域，避免删除原结构；仅居中最初页脚段落。',f'第{j+1}节页脚','核对自动页码及连续性')

    def review_numbering_and_sources(self):
        """Structural lint: a warning is not an authorization to rewrite citations."""
        stack={}
        for h in self.headings:
            text=h['text']
            level=h['level']
            m=re.match(r'^([0-9]{1,3}(?:\.[0-9]{1,3}){0,3})\s*[.．、]?\s*',text)
            if m:
                nums=m.group(1).split('.')
                if len(nums)>=2:
                    parent='.'.join(nums[:-1])
                    if stack.get(level-1)!=parent:
                        self.issue('review','heading_hierarchy',
                            f'标题层级可能缺父级或与前面的编号不连续：{text[:70]}（应有父级 {parent}）。',
                            f'正文第{h["paragraph"]}段','对照正文标题树复核章节序号')
                stack[level]=m.group(1)
                for lower in range(level+1,6):stack.pop(lower,None)
        for caption in self.captions:
            i=caption['paragraph']-1
            nearby=' '.join(p.text for p in self.doc.paragraphs[max(0,i-2):min(len(self.doc.paragraphs),i+5)])
            if not re.search(r'资料来源|数据来源|图片来源|Source\s*[:：]|Sources\s*[:：]',nearby,re.I):
                self.issue('review','missing_source',
                           f'未在图表题附近识别到明确资料来源：{caption["text"][:75]}；也可能来源嵌在图片、表内或脚注。',
                           f'图表题第{caption["paragraph"]}段','补齐可核实的真实来源，禁止虚构来源')
        table_caption_num=sum(x['kind']=='table_caption' for x in self.captions)
        if self.doc.tables and table_caption_num<len(self.doc.tables):
            self.issue('review','table_caption_count',
                       f'检测到 {len(self.doc.tables)} 张 Word 表格，但只发现 {table_caption_num} 个常规表题；特殊标题或无标题表格需要人工核查。',
                       '全文表格','按章编号并在表格正上方放置宋体小五粗体表题')
        stats=package_xml_stats(self.infile)
        figure_captions=sum(x['kind']=='figure_caption' for x in self.captions)
        if stats['pictures'] and figure_captions==0:
            self.issue('review','missing_figure_caption',
                       f'检测到 {stats["pictures"]} 个主文档绘图元素，却没有识别到“图 1-1 ...”形式的图题。',
                       '全文图片','逐一补齐图题、按章编号及来源')
        if stats['math_objects']:
            self.issue('review','math_layout',
                       f'检测到 {stats["math_objects"]} 个 Word 数学对象，已保留原始公式；公式另行、缩格、编号右对齐及跨行换行位置无法靠本引擎保证。',
                       '公式','逐个在 Word 中验证，核对每章公式编号')
        if stats['has_toc_field']:
            self.issue('review','dynamic_toc','检测到 Word 目录域，格式化及分页完成后需更新整张目录（包括页码）。',
                       '目录','Word 中选中目录 → 更新整个目录')
        elif self.markers.get('toc_title'):
            self.issue('review','static_toc','检测到目录标题，但未识别到 Word 的 TOC 域；目录可能是静态文字，页码不会自动刷新。',
                       '目录','使用 Word 基于真实标题样式重建动态目录并更新页码')

    def review_caption_sequences(self):
        # Derive the chapter from an explicit caption and review duplicates,
        # non-sequential numbers and conflicts with currently identified chapter.
        found=defaultdict(list)
        heading_at=[]
        for h in self.headings:
            if h['level']==1:
                m=re.match(r'^\s*(\d{1,3})(?:[.．、\s\u3000]|$)',h['text'])
                if m:heading_at.append((h['paragraph'],int(m.group(1))))
        for c in self.captions:
            match=re.match(r'^\s*([表图])\s*(\d{1,3})[-－—](\d{1,3})(?=\s|\S)',c['text'])
            if not match:
                self.issue('review','caption_number_pattern',
                     '图表题不是正式规范所举的“图 3-1 / 表 2-7”编号形式，保持原文：'+c['text'][:65],
                     f'第{c["paragraph"]}段','检查编号及 Word 交叉引用')
                continue
            kind,chapter,ordinal=match.group(1),int(match.group(2)),int(match.group(3))
            found[(kind,chapter)].append((ordinal,c['paragraph']))
            prior=[number for pos,number in heading_at if pos<c['paragraph']]
            if prior and prior[-1]!=chapter:
                self.issue('review','caption_chapter_mismatch',
                     f'{kind}题编号归于第{chapter}章，但前面最近的一级标题为第{prior[-1]}章。',
                     f'第{c["paragraph"]}段','核对图表编号、位置及交叉引用')
        for (kind,chapter),numbers in found.items():
            seq=[x[0] for x in numbers]
            if len(seq)!=len(set(seq)):
                self.issue('review','caption_duplicate',
                    f'第{chapter}章{kind}编号重复：{seq}。',f'第{chapter}章','人工核对重复编号及正文交叉引用')
            actual=sorted(set(seq)); expected=list(range(1,max(actual)+1))
            if actual!=expected:
                self.issue('review','caption_sequence',
                    f'第{chapter}章{kind}编号存在缺号：{actual}，预计连续序号为{expected}。',
                    f'第{chapter}章','检查是否有图题嵌于文本框/跨章节引用，确认后统一')
        self.diagnostics['caption_groups_checked']=len(found)

    def supplemental_checks(self):
        if not self.markers.get('abstract_cn'):
            self.issue('review','missing_component','未识别到独立中文“摘要”标题。','论文前置内容','核对中英文摘要及关键词')
        if not self.markers.get('abstract_en'):
            self.issue('review','missing_component','未识别到独立英文 Abstract 标题。','论文前置内容','核对中英文摘要及分页')
        if not self.markers.get('toc_title'):
            self.issue('review','missing_component','未识别到独立“目录”标题。','论文前置内容','核对目录及三级标题')
        if not self.markers.get('references_title'):
            self.issue('review','missing_component','未识别到独立“参考文献”标题。','论文末尾','核对参考文献位置与引用规范')
        if self.markers.get('abstract_cn') and self.markers.get('abstract_en'):
            cn_i=self.markers['abstract_cn'][0];en_i=self.markers['abstract_en'][0]
            if cn_i<en_i:
                cn=''.join(p.text for p in self.doc.paragraphs[cn_i+1:en_i]
                           if not re.search(r'关键词|key ?words',p.text,re.I))
                count=len(re.findall(r'[\u4e00-\u9fff]',cn))
                if not 300<=count<=500:
                    self.issue('info','abstract_length',f'中文摘要约含{count}个汉字，学校建议约300—500字；数量仅为粗略统计。',
                               '中文摘要','必要时作者自行精简/补充内容')
        refs=self.markers.get('references_title',[])
        if refs:
            start=refs[0]+1
            stop=min([a for k in ('appendix_title','ack_title') for a in self.markers.get(k,[]) if a>start] or [len(self.doc.paragraphs)])
            raw=[p.text.strip() for p in self.doc.paragraphs[start:stop] if p.text.strip()]
            count=sum(bool(re.match(r'^\s*\[?\d+\]?\s*[.．、]?\s*\S+',x)) for x in raw)
            if count<30:
                self.issue('info','references_count',f'自动数到约{count}条编号参考文献（格式可能使计数低估），校规为一般不少于30篇/本。',
                           '参考文献','人工核查文献数及外文原文')
            self.issue('info','references_check','已保留原引文和著录内容；GB/T 7714—2015 格式一致性、外文原文及正文引用对应关系仍需人工核对。',
                       '参考文献','不可将基本样式修正等同于文献事实校验')
        if self.markers.get('ack_title'):
            ack_idx=self.markers['ack_title'][0]
            txt=''.join(p.text for p in self.doc.paragraphs[ack_idx+1:])
            c=len(re.findall(r'[\u4e00-\u9fff]',txt))
            if c>1000:self.issue('review','ack_length',f'致谢约{c}个汉字，校规一般不超过1000个汉字且限一页。','致谢','人工删减，并视觉检查是否超过一页')
        self.issue('review','mandatory_visual_qa','字号与边距可机器审查，但尚未验证所有分页、悬浮图锚点、公式换行、续表及奇数页起始。','全稿','在 Word/LibreOffice 渲染全页并逐页核对')
        self.issue('info','cover','封面、独创性声明和授权声明须采用研究生院模板且有本人/导师亲笔签字；程序无法代替签字或保证现有封面为官方模板。','前置部分','根据研究生院模板单独检查')
        self.issue('info','source_conflict','正式规范“第一层：1.”，配套范例正文使用“1 标题”，本工具不改变标题编号字符本身。','一级标题','如学校提供更新版模板，以更新版为准')
        self.issue('info','unspecified_rule','正式规范写“章的标题占2行”，但范例标题显示并非逐章强制硬换行；本工具用“与下段同页＋整章另页”，不人工插入空行。','一级标题','毕业定稿前与指导教师/学院核对')

    def run(self):
        self.set_section_layout()
        self.set_styles()
        self.process_paragraphs()
        self.process_tables()
        self.check_caption_locations()
        self.ensure_body_section()
        self.footnotes()
        self.format_headers_footers()
        self.review_numbering_and_sources()
        self.review_caption_sequences()
        self.supplemental_checks()


def normalize_footnote_xml(path):
    """Apply 9pt, single-spacing in existing footnote XML in-place via ZIP rewrite."""
    from tempfile import NamedTemporaryFile
    import os
    with ZipFile(path) as z:
        if 'word/footnotes.xml' not in z.namelist():return 0
        xml=z.read('word/footnotes.xml')
    root=etree.fromstring(xml)
    n=0
    for fn in root.xpath('./w:footnote', namespaces=NS):
        if fn.get(qn('w:type')) in ('separator','continuationSeparator'):continue
        for p in fn.xpath('./w:p',namespaces=NS):
            ppr=p.find(qn('w:pPr'))
            if ppr is None:
                ppr=OxmlElement('w:pPr');p.insert(0,ppr)
            spacing=ppr.find(qn('w:spacing'))
            if spacing is None:
                spacing=OxmlElement('w:spacing');ppr.append(spacing)
            spacing.set(qn('w:line'),'240');spacing.set(qn('w:lineRule'),'auto')
            for r in p.xpath('.//w:r',namespaces=NS):
                if r.xpath('./w:fldChar|./w:instrText|./w:drawing|./w:object|./m:oMath',namespaces=NS):continue
                rpr=r.find(qn('w:rPr'))
                if rpr is None:
                    rpr=OxmlElement('w:rPr');r.insert(0,rpr)
                for tag in ('sz','szCs'):
                    sz=rpr.find(qn('w:'+tag))
                    if sz is None:sz=OxmlElement('w:'+tag);rpr.append(sz)
                    sz.set(qn('w:val'),'18')
                n+=1
    with NamedTemporaryFile(suffix='.docx',delete=False,dir=str(path.parent)) as temp:
        tmp=temp.name
    try:
        with ZipFile(path) as src, ZipFile(tmp,'w') as dst:
            for info in src.infolist():
                data=etree.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True) if info.filename=='word/footnotes.xml' else src.read(info.filename)
                dst.writestr(info,data)
        os.replace(tmp,path)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
    return n


def assess_integrity(inp,out,before,after):
    notes=[]
    for key in ('paragraphs','tables'):
        if before[key]!=after[key]:notes.append({'severity':'severe','detail':f'{key}数量变化: {before[key]} → {after[key]}'})
    for key in ('paragraph_texts','table_texts'):
        if before[key]!=after[key]:
            canon=lambda cnt:Counter({re.sub(r'[\s\u3000]+','',k):sum(vv for kk,vv in cnt.items() if re.sub(r'[\s\u3000]+','',kk)==re.sub(r'[\s\u3000]+','',k)) for k in cnt})
            if key=='paragraph_texts' and canon(before[key])==canon(after[key]):
                notes.append({'severity':'ok','detail':'个别标题文字仅调整了编号后的空白间距；去空白字符后全部内容一致。'})
            else:
                notes.append({'severity':'severe','detail':f'全文{key}集合变化，可能发生实质文字损坏。'})
    orig=media_fingerprint(inp); new=media_fingerprint(out)
    if orig!=new:notes.append({'severity':'severe','detail':'word/media 文件或哈希变化，图片/媒体完整性存疑。'})
    # Guard independent binary/XML parts python-docx should not rewrite: charts,
    # Excel links/embedded sheets, custom diagrams, comments, glossary, embeddings.
    guarded_prefixes=('word/charts/','word/embeddings/','word/diagrams/',
                      'word/activeX/','customXml/','word/comments')
    def guarded(path):
        with ZipFile(path) as z:
            return {n: hashlib.sha256(z.read(n)).hexdigest()
                    for n in z.namelist() if n.startswith(guarded_prefixes)}
    a,b=guarded(inp),guarded(out)
    if a!=b:
        lost=sorted(set(a)-set(b)); changed=sorted(k for k in a.keys() & b.keys() if a[k]!=b[k])
        added=sorted(set(b)-set(a))
        notes.append({'severity':'severe',
            'detail':f'附属对象包可能丢失/改变：丢失{lost[:10]}、变化{changed[:10]}、新增{added[:10]}。'})
    with ZipFile(out) as z:
        bad=z.testzip()
        if bad:notes.append({'severity':'severe','detail':'ZIP 内成员CRC损坏: '+bad})
    for k in ('math_objects','pictures','tracked_change_nodes','has_comments'):
        a=package_xml_stats(inp)[k];b=package_xml_stats(out)[k]
        if a!=b:notes.append({'severity':'severe','detail':f'{k}: {a} → {b}'})
    return notes or [{'severity':'ok','detail':'段落/表格数量与文字、媒体 SHA256、主要公式/图片/修订/批注统计均未见意外变化；ZIP CRC 检查通过。'}]


def report_markdown(result):
    summary=result['summary']
    out=[f'# {summary["input_file"]} — 格式修正报告', '',
         f'- **规范基准**：{NAME} + 校方格式范例；非自动证明当年最新。',
         f'- **输入文件**：`{summary["input_file"]}`',
         f'- **输出文件**：`{summary["output_file"]}`',
         f'- **执行时间**：{summary["time"]}',
         f'- **工具版本**：{VERSION}（GitHub 开源模板对照审查升级版）',
         f'- **标题识别结构**：{summary["number_mode"]}；题目：{summary["title"] or "未可靠识别"}',
         f'- **原稿规模**：{summary["paragraphs"]}段，{summary["tables"]}张表；含{summary["pictures"]}个绘图元素、{summary["math_objects"]}个数学公式对象。',
         f'- **合计更改事件**：{len(result["changes"])}；待复核/提示项：{len(result["issues"])}。',
         f'- **语义底色表保护**：{summary.get("diagnostics",{}).get("table_semantic_colors_preserved",0)} 张；**安全移动图题**：{summary.get("diagnostics",{}).get("figure_caption_auto_relocated",0)} 处；**已核查图表编号分组**：{summary.get("diagnostics",{}).get("caption_groups_checked",0)} 组。',
         '', '## 1. 格式修改统计', '', '| 类型 | 数量 |', '|---|---:|']
    for k,n in sorted(summary['changes_by_type'].items()):out.append(f'| {k} | {n} |')
    out+=['','## 2. 自动识别的标题层级','','| 段落编号 | 级别 | 置信度 | 识别内容 |', '|---:|---:|---|---|']
    for x in result['headings'][:180]:
        out.append(f'| {x["paragraph"]} | {x["level"]} | {x["confidence"]} | {x["text"].replace("|","/")[:90]} |')
    if len(result['headings'])>180:out.append(f'\n> 共 {len(result["headings"])} 个标题；剩余列表见 JSON 报告。')
    out+=['','## 3. 待人工复核 / 歧义与风险','','> 请先查看“严重”和“需要复核”事项。机器无法验证的项目明确列出，不代表全部已通过。','']
    for sev, title in [('severe','严重'),('review','需要复核'),('info','说明/非硬性条款')]:
        bucket=[x for x in result['issues'] if x['severity']==sev]
        out.append(f'### {title}（{len(bucket)}项）')
        for x in bucket[:120]:
            out.append(f'- **[{x["category"]}]** {x["location"]}：{x["detail"]} 处理：{x["recommended_action"]}。')
        if len(bucket)>120:out.append(f'- 其余 {len(bucket)-120} 项见 JSON 报告。')
        out.append('')
    out+=['## 4. 内容和对象完整性校验','']
    for item in result['integrity']:
        out.append(f'- **{item["severity"]}**：{item["detail"]}')
    out+=['','## 5. 前后格式变更明细（前120项）','','| 对象 | 位置 | 变更前 | 变更后 | 校方依据 |', '|---|---|---|---|---|']
    for c in result['changes'][:120]:
        clean=lambda s:str(s).replace('|','/').replace('\n',' ')[:100]
        out.append(f'| {clean(c["target"])} | {clean(c["location"])} | {clean(c["before"])} | {clean(c["after"])} | {clean(c["evidence"])} |')
    if len(result['changes'])>120:out.append(f'\n> 其余 {len(result["changes"])-120} 项记录在 JSON 报告。')
    out+=['','## 6. 最终交付前的 Word/视觉核查','','- [ ] Word 能正常打开、没有断图/乱码/布局溢出。',
          '- [ ] 逐页审查：每一章另起页、正文右页起始、图题在下/表题在上。',
          '- [ ] 页眉为论文真实题目，正文页码居中，封面页码正确隔离，多节页码无错误。',
          '- [ ] 目录重新更新字段并核对至三级、真实页码与标题一致。',
          '- [ ] 图表编号按章连续、来源真实完整，复杂表格与跨页“续表”处理无误。',
          '- [ ] 参考文献国标格式、脚注编号、外文文献及引用关系人工核验。',
          '- [ ] 官方封面、授权声明、签名均通过人工核查。',
          '- [ ] 2020版要求是否仍为适用版本，已与当年院系要求核实。',
          '', '## 7. 无法自动保证的范围','','扫描 PDF 的无损 Word 化、浮动对象最终视觉位置、书脊及亲笔签名、Word 字段刷新与打印端差异均需相应的人工或办公软件处理。', '']
    return '\n'.join(out)


def main(argv=None):
    pa=argparse.ArgumentParser(description='Format a .docx thesis to SUIBE 2020 master thesis specification, generate audit reports.')
    pa.add_argument('input',type=Path,help='editable DOCX source (not scanned PDF)')
    pa.add_argument('--out',type=Path,help='output DOCX')
    pa.add_argument('--report',type=Path,help='Markdown report')
    pa.add_argument('--report-json',type=Path,help='JSON report')
    pa.add_argument('--title',help='actual dissertation title, for page header')
    pa.add_argument('--table-mode',choices=['simple','preserve'],default='simple',help='simple: only ordinary unmerged data tables to white three-line; preserve: do not touch table borders')
    pa.add_argument('--headers',choices=['safe','preserve'],default='safe')
    args=pa.parse_args(argv)
    inp=args.input.resolve()
    if not inp.exists():pa.error('input does not exist')
    if inp.suffix.lower()!='.docx':pa.error('only .docx is accepted. Convert legacy .doc first with layout QA; scanned PDFs are not editable Word.')
    if args.out:
        out=args.out.resolve()
    else:out=inp.with_name(inp.stem+'_SUIBE规范排版.docx')
    if out==inp:pa.error('will not overwrite input file')
    rpt=(args.report or out.with_name(inp.stem+'_格式修正报告.md')).resolve()
    rptjson=(args.report_json or out.with_name(inp.stem+'_格式修正报告.json')).resolve()
    for path in [out,rpt,rptjson]:path.parent.mkdir(parents=True,exist_ok=True)
    try:
        doc=Document(str(inp))
        before=simple_structure(doc)
        st=package_xml_stats(inp)
        work=Work(doc,inp,title=args.title,table_mode=args.table_mode,headers=args.headers)
        if st['tracked_change_nodes']:
            work.issue('review','tracked_changes',f'检测到 {st["tracked_change_nodes"]} 个修订节点；没有擅自接受/拒绝修订，相关复杂段落跳过修改。','全文','用 Word 人工处理修订并再次核查')
        if st['complex_control_count']:
            work.issue('review','complex_objects',f'检测到 {st["complex_control_count"]} 个内容控件/文本框/OLE对象（检测值含嵌套）；可能需要额外 OOXML 处理。','全文','视觉检查嵌套对象')
        if st['has_comments']:
            work.issue('info','comments','检测到 Word 批注关系；程序不会删除批注。','全文')
        work.run()
        doc.save(str(out))
        nf=normalize_footnote_xml(out)
        if nf:work.change('脚注',f'脚注文字run {nf}段','混合/旧格式','小五9pt、单倍行距，Word脚注按页重启编号','规范第4页')
        reread=Document(str(out))
        integrity=assess_integrity(inp,out,before,simple_structure(reread))
        for x in integrity:
            if x['severity']=='severe':work.issue('severe','data_integrity',x['detail'],'整体','禁止作为定稿交付，人工修复')
        summary={'input_file':inp.name,'output_file':out.name,'time':datetime.now().isoformat(timespec='seconds'),
                 'version':VERSION,'number_mode':work.mode,'title':work.title,'paragraphs':before['paragraphs'],
                 'tables':before['tables'],'pictures':st['pictures'],'math_objects':st['math_objects'],
                 'changes_by_type':dict(work.counts),'diagnostics':work.diagnostics}
        result={'summary':summary,'headings':work.headings,'captions':work.captions,
                'issues':work.issues,'changes':work.changes,'integrity':integrity,
                'classification':work.recognized}
        rpt.write_text(report_markdown(result),encoding='utf-8')
        rptjson.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        print('输出Word:',out)
        print('报告Markdown:',rpt)
        print('报告JSON:',rptjson)
        print('识别标题:',len(work.headings),'|更改事件:',len(work.changes),'|待复核/信息:',len(work.issues))
        return 2 if any(i['severity']=='severe' for i in work.issues) else 0
    except Exception:
        traceback.print_exc()
        return 1

if __name__=='__main__':sys.exit(main())
